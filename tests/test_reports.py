"""Report boundaries and persistence; isolated data, no external API calls."""
import importlib.util
import io
import os
import sys
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from fastapi import HTTPException
from fastapi.testclient import TestClient

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'backend'))


class ReportsTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.env=patch.dict(os.environ,{'TICHA_DATA_DIR':self.temp.name,'TICHA_AI_KEY':'','TICHA_AI_MODEL':''})
        self.env.start()
        spec=importlib.util.spec_from_file_location('report_test_app',ROOT/'backend/app.py')
        self.module=importlib.util.module_from_spec(spec);spec.loader.exec_module(self.module)
        self.client=TestClient(self.module.app)
        self.pid=self.client.post('/api/projects',json={'name':'测试项目'}).json()['id']
        self.file=self.client.post(f'/api/projects/{self.pid}/files',files={'file':('现场记录.txt','王店主：装卸通道经常被占用。'.encode(),'text/plain')}).json()
        self.sources={};self.templates=set()

    def tearDown(self):
        self.client.close();self.env.stop();self.temp.cleanup()

    def tool(self,name,args,choice=None):
        return self.module.agent_tool(self.pid,name,args,'请根据材料生成报告',self.sources,self.templates,choice)

    def test_report_requires_read_and_stays_in_project(self):
        args={'name':'需求报告','content':'# 需求报告\n\n装卸通道占用，见[现场记录.txt:P1]。','source_ids':[self.file['id']],'template_id':'none'}
        with self.assertRaises(HTTPException):self.tool('create_report',args,'none')
        self.tool('read_document',{'file_id':self.file['id']})
        result=self.tool('create_report',args,'none')
        rows=self.client.get(f'/api/projects/{self.pid}/files').json()
        report=next(row for row in rows if row['id']==result['id'])
        self.assertEqual(report['kind'],'report')
        saved=self.client.get(f'/api/projects/{self.pid}/files/{result["id"]}/content').json()
        self.assertIn('现场记录.txt',saved['content']);self.assertIn('本报告材料索引',saved['content'])
        other=self.client.post('/api/projects',json={'name':'另一项目'}).json()['id']
        self.assertEqual(self.client.get(f'/api/projects/{other}/files/{result["id"]}/content').status_code,404)

    def test_template_missing_sections_and_wrong_choice_rejected(self):
        tid='market-research'
        self.tool('read_document',{'file_id':self.file['id']})
        self.tool('read_report_template',{'template_id':tid})
        args={'name':'市场报告','content':'# 不完整的报告','source_ids':[self.file['id']],'template_id':tid}
        with self.assertRaises(HTTPException):self.tool('create_report',args,'none')
        with self.assertRaises(HTTPException):self.tool('create_report',args,tid)
        args['content']='# 市场报告\n\n'+'\n\n'.join('#'*(s['level']+1)+' '+s['title']+'\n\n本项材料未提供，待补充。' for s in self.module.report_templates.TEMPLATES_BY_ID[tid]['sections'])
        result=self.tool('create_report',args,tid)
        self.assertTrue(result['saved']);self.assertEqual(result['template_id'],tid)

    def test_material_change_and_stale_editor_save_rejected(self):
        fid=self.file['id']
        read=self.tool('read_document',{'file_id':fid})
        response=self.client.put(f'/api/projects/{self.pid}/files/{fid}/content',json={'name':'现场记录.txt','content':'新增证据','expected_hash':read['hash']})
        self.assertEqual(response.status_code,200);self.assertIn('hash',response.json())
        stale=self.client.put(f'/api/projects/{self.pid}/files/{fid}/content',json={'name':'现场记录.txt','content':'旧编辑器覆盖','expected_hash':read['hash']})
        self.assertEqual(stale.status_code,409)
        with self.assertRaises(HTTPException):self.tool('create_report',{'name':'报告','content':'# 报告','source_ids':[fid],'template_id':'none'},'none')

    def test_docx_extracts_tables_and_paragraphs(self):
        from docx import Document
        document=Document();document.add_paragraph('访谈事实');table=document.add_table(rows=1,cols=2);table.cell(0,0).text='装卸';table.cell(0,1).text='通道占用'
        buffer=io.BytesIO();document.save(buffer)
        uploaded=self.client.post(f'/api/projects/{self.pid}/files',files={'file':('访谈.docx',buffer.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()
        result=self.tool('read_document',{'file_id':uploaded['id']})
        self.assertIn('访谈事实',result['content']);self.assertIn('通道占用',result['content']);self.assertTrue(result['paragraphs'])

if __name__=='__main__':unittest.main()
